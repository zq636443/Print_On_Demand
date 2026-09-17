from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT='中国版POD商品创作与履约中台_PRD.docx'
doc=Document(); sec=doc.sections[0]; sec.top_margin=Inches(.65); sec.bottom_margin=Inches(.65); sec.left_margin=Inches(.72); sec.right_margin=Inches(.72)
styles=doc.styles
styles['Normal'].font.name='Heiti SC'; styles['Normal']._element.rPr.rFonts.set(qn('w:eastAsia'),'Heiti SC'); styles['Normal'].font.size=Pt(10.5)
for s,size in [('Title',24),('Heading 1',16),('Heading 2',12.5),('Heading 3',11)]:
    styles[s].font.name='Heiti SC'; styles[s]._element.rPr.rFonts.set(qn('w:eastAsia'),'Heiti SC'); styles[s].font.size=Pt(size); styles[s].font.color.rgb=RGBColor(0,0,0)
def shade(cell, fill):
    tcPr=cell._tc.get_or_add_tcPr(); shd=OxmlElement('w:shd'); shd.set(qn('w:fill'),fill); tcPr.append(shd)
def borders(cell):
    tcPr=cell._tc.get_or_add_tcPr(); b=OxmlElement('w:tcBorders')
    for e in ['top','left','bottom','right','insideH','insideV']:
        x=OxmlElement('w:'+e); x.set(qn('w:val'),'single'); x.set(qn('w:sz'),'4'); x.set(qn('w:color'),'D9D9D9'); b.append(x)
    tcPr.append(b)
def table(headers, rows, widths=None):
    t=doc.add_table(rows=1, cols=len(headers)); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.style='Table Grid'
    for i,h in enumerate(headers):
        c=t.rows[0].cells[i]; c.text=h; shade(c,'1F4E78'); borders(c); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for r in c.paragraphs[0].runs: r.font.bold=True; r.font.color.rgb=RGBColor(255,255,255)
    for ri,row in enumerate(rows):
        cells=t.add_row().cells
        for i,v in enumerate(row):
            cells[i].text=str(v); borders(cells[i]); cells[i].vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if ri%2: shade(cells[i],'F3F6F9')
    doc.add_paragraph(''); return t
def bullets(items):
    for x in items: doc.add_paragraph(x, style='List Bullet')

p=doc.add_paragraph(style='Title'); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run('中国版 POD 商品创作与履约中台 PRD')
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run('讨论稿 v0.1  |  面向创作者、商家与供应商').italic=True
doc.add_paragraph('本产品不是独立电商商城，而是连接商品创作端、电商销售端与生产供应端的中台。创作者在本平台完成设计并授权自己的淘宝、抖音或小红书店铺发布商品；买家始终在原电商平台付款；订单回传后，本平台负责 SKU 映射、供应商分发、生产状态与发货信息回传。本文用于产品、技术和业务讨论，重点明确功能边界与首期取舍。')

doc.add_heading('一 产品定位与目标',1)
doc.add_paragraph('定位：商品创作工具 + 多渠道商品管理 + 订单履约中台。')
table(['目标','说明'],[['创作者低门槛商品化','上传设计，生成可售卖商品并发布到自有店铺'],['销售端不迁移','交易、收款、买家关系仍属于淘宝/抖音/小红书店铺'],['供应端标准化履约','将订单转化为标准生产单，分发给合适供应商'],['平台可持续变现','设计增值服务、订单技术服务费、供应商服务费等']])
doc.add_heading('二 角色与权限',1)
table(['角色','核心权限'],[['创作者/商家','设计商品、绑定店铺、发布商品、查看订单与收益'],['供应商','维护产能与 SKU、接单、更新生产状态、填写发货信息、处理售后'],['平台运营','审核用户和商品、管理供应商、处理异常与结算'],['系统管理员','渠道授权、接口配置、权限、日志和安全管理']])
doc.add_heading('三 总体业务链路',1)
doc.add_paragraph('创作者设计 → 授权销售渠道 → 发布/更新渠道商品 → 买家在渠道下单 → 渠道订单同步 → 平台生成履约单 → 供应商接单生产 → 回传运单 → 平台回传渠道发货状态 → 渠道向买家展示物流 → 订单完成与结算。')

doc.add_heading('四 功能模块',1)
modules=[('4.1 创作中心',['素材上传与管理：PNG/JPG/WebP、分辨率检测、版权声明。','商品设计器：选择模板、颜色/尺码、图片缩放旋转、文字、印刷区域、Mockup。','设计版本：保存草稿、复制、回滚、关联商品。','内容审核：机器初筛 + 人工审核 + 投诉下架。']),('4.2 商品中心',['商品模板与 SKU：商品、颜色、尺码、印刷工艺、供应商 SKU 的标准映射。','定价：生产成本、平台服务费、创作者预期收益、渠道费用、售价。','商品状态：草稿、待审核、已通过、发布中、已发布、同步失败、已下架。','渠道映射：平台商品 ID ↔ 渠道商品 ID ↔ 供应商 SKU。']),('4.3 渠道中心',['店铺授权：淘宝、抖音、小红书等，保存授权状态、店铺名称和权限范围。','商品发布：自动发布、生成素材包、半自动辅助三种模式。','商品同步：标题、图片、详情、SKU、价格、库存、发货时效。','发布结果：成功、失败、待人工处理，并记录错误原因。']),('4.4 订单履约中心',['订单归集：同步渠道订单、付款状态、收货信息、买家备注。','订单标准化：渠道订单转内部履约订单。','生产分发：按商品绑定、区域、产能、时效和供应商状态路由。','状态机：待生产、待接单、生产中、待发货、已发货、已完成、异常。','异常处理：拒单、缺货、超时、地址错误、生产失败、重复订单。']),('4.5 供应商中心',['入驻与认证：营业执照、工厂信息、品类、产能、生产时效。','供应商 SKU：成本、工艺、可生产颜色/尺码、最低订单量。','接单工作台：查看生产文件、规格、地址，接受/拒绝/异常。','发货回填：快递公司、运单号、发货时间、凭证。','评分：准时率、质量投诉率、拒单率、售后率。']),('4.6 售后与风控',['质量问题、错发漏发、物流异常、买家选错规格等分类处理。','创作者版权承诺、敏感内容审核、侵权投诉、证据留存。','超时预警、自动冻结结算、供应商降权或暂停接单。']),('4.7 结算中心',['平台服务费、创作者收益、供应商生产费用、退款和预留金。','账单、对账、提现申请、结算周期、发票信息。','首期建议按订单完成后 T+7/T+15 生成可结算金额。']),('4.8 运营后台',['用户/店铺/供应商管理、商品审核、订单监控、异常工单、财务对账、操作日志、数据看板。'])]
for title, items in modules:
    doc.add_heading(title,2); bullets(items)

doc.add_heading('五 核心数据对象',1)
table(['对象','关键字段','关系'],[['平台商品','商品 ID、设计文件、模板、售价、生产方案','关联多个渠道商品与供应商 SKU'],['渠道商品','渠道、店铺、渠道商品 ID、同步状态','对应一个平台商品'],['渠道订单','渠道订单 ID、店铺、买家地址、订单状态','生成一个或多个履约单'],['履约订单','内部单号、平台商品、SKU、生产参数、供应商','驱动生产与发货'],['物流记录','快递公司、运单号、发货时间、轨迹状态','回传渠道并供查询']])
doc.add_heading('六 首期 MVP 建议',1)
bullets(['优先支持淘宝，抖音作为第二渠道；小红书先支持素材包或人工辅助发布，待确认官方权限后再自动化。','创作者可授权店铺，但商品发布首期允许“手动发布 + 订单导入”作为兜底。','供应商使用网页后台接单和回填运单，暂不要求 ERP/API 对接。','物流首期只做运单回传给电商平台，不自建物流查询系统。','平台不承接买家支付，先通过技术服务费、供应商服务费或生产差价变现。'])
doc.add_heading('七 关键接口与依赖',1)
table(['接口方向','用途','首期策略'],[['电商平台 → 本平台','订单、退款、商品状态通知','授权后轮询 + webhook（以平台开放能力为准）'],['本平台 → 电商平台','创建/更新商品、回传发货','优先淘宝，抖音后续，小红书按权限确认'],['本平台 → 供应商','生产单、设计文件、规格和地址','供应商后台；API 作为后续能力'],['供应商 → 本平台','接单、生产状态、运单号','后台回填；支持批量导入'],['本平台 → 物流服务','轨迹查询','MVP 不做，后续接第三方 API']])
doc.add_heading('八 非功能要求',1)
bullets(['安全：店铺授权凭证加密存储；收货地址最小化使用；操作日志可追溯。','可靠性：订单幂等，避免重复生产；接口失败可重试；异常订单可人工补偿。','时效：订单同步延迟目标小于 5 分钟；发货回传失败可告警。','可运营：所有自动化动作保留人工兜底入口。','合规：版权、隐私、平台规则、个性化商品售后规则在上线前完成审核。'])
doc.add_heading('九 待讨论决策',1)
table(['问题','建议讨论选项'],[['首个正式渠道','淘宝优先 / 抖音优先 / 先做素材包'],['商品发布方式','全自动 / 半自动 / 手动兜底'],['订单导入方式','API / webhook / 轮询 / 文件导入'],['供应商路由','商品绑定 / 价格优先 / 时效优先 / 区域优先'],['物流范围','只回传运单 / 同步轨迹 / 电子面单'],['收费方式','按单服务费 / 月费 / 供应商服务费 / 差价'],['售后责任','平台统一规则 / 供应商承担质量责任 / 按类目配置']])
doc.add_heading('十 一句话版本',1)
doc.add_paragraph('平台负责把“设计”变成“渠道商品”，把“渠道订单”变成“供应商生产单”，再把“供应商运单”回传给原电商平台；买家交易始终发生在淘宝、抖音或小红书，而不是本平台。')
doc.save(OUT)
print(OUT)
